"""Render a document source into a file.

The governing idea, from the owner: **a document is its source, not its
bytes.** A PDF cannot be edited meaningfully, but the Markdown that produced
it can — so "make the heading bigger" is a new render of an edited source,
not a byte patch. Every format here is therefore a pure function from a
stored source to a file, and re-rendering is always possible.

Backends are optional. A format whose library is missing is reported as
unavailable with the exact install command, rather than failing somewhere
deep in a traceback.
"""

from __future__ import annotations

import csv
import html as html_escape
import io
import json
import logging
import re
from dataclasses import dataclass
from pathlib import Path

logger = logging.getLogger(__name__)

DOCUMENT = "document"
SHEET = "sheet"
SLIDES = "slides"

#: format -> (kind it belongs to, pip package needed, install hint)
FORMATS: dict[str, tuple[str, str | None, str]] = {
    "md": (DOCUMENT, None, ""),
    "txt": (DOCUMENT, None, ""),
    "html": (DOCUMENT, None, ""),
    "pdf": (DOCUMENT, "reportlab", "pip install reportlab"),
    "csv": (SHEET, None, ""),
    "json": (SHEET, None, ""),
    "xlsx": (SHEET, "openpyxl", "pip install openpyxl"),
    "pptx": (SLIDES, "pptx", "pip install python-pptx (needs pkg install python-lxml python-pillow first)"),
}


class RenderError(RuntimeError):
    def __init__(self, message: str, *, code: str = "render_failed"):
        super().__init__(message)
        self.code = code


@dataclass(frozen=True)
class Block:
    """One parsed piece of a Markdown source."""

    kind: str  # heading | paragraph | bullet | number | code
    text: str
    level: int = 0


_HEADING = re.compile(r"^(#{1,6})\s+(.*)$")
_BULLET = re.compile(r"^\s*[-*+]\s+(.*)$")
_NUMBER = re.compile(r"^\s*\d+[.)]\s+(.*)$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ITALIC = re.compile(r"(?<!\*)\*([^*\n]+)\*(?!\*)")
_CODE = re.compile(r"`([^`\n]+)`")


def parse_markdown(source: str) -> list[Block]:
    """A deliberately small Markdown subset.

    Small because every construct here must render in four different
    backends. Anything exotic would work in HTML and silently vanish in PDF,
    which is worse than not supporting it.
    """
    blocks: list[Block] = []
    paragraph: list[str] = []
    fenced: list[str] | None = None

    def flush() -> None:
        if paragraph:
            blocks.append(Block("paragraph", " ".join(paragraph).strip()))
            paragraph.clear()

    for raw in (source or "").replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if line.strip().startswith("```"):
            if fenced is None:
                flush()
                fenced = []
            else:
                blocks.append(Block("code", "\n".join(fenced)))
                fenced = None
            continue
        if fenced is not None:
            fenced.append(raw)
            continue
        if not line.strip():
            flush()
            continue
        heading = _HEADING.match(line)
        if heading:
            flush()
            blocks.append(Block("heading", heading.group(2).strip(), len(heading.group(1))))
            continue
        bullet = _BULLET.match(line)
        if bullet:
            flush()
            blocks.append(Block("bullet", bullet.group(1).strip()))
            continue
        number = _NUMBER.match(line)
        if number:
            flush()
            blocks.append(Block("number", number.group(1).strip()))
            continue
        paragraph.append(line.strip())

    flush()
    if fenced is not None:  # unterminated fence: keep the content rather than lose it
        blocks.append(Block("code", "\n".join(fenced)))
    return blocks


def drop_duplicate_title(blocks: list[Block], title: str) -> list[Block]:
    """Remove a leading heading that just repeats the title.

    Models reliably open a Markdown document with `# Title`, and every
    renderer here also prints the title. Without this the output says it
    twice, which looks like a bug to the reader.
    """
    if blocks and blocks[0].kind == "heading":
        if strip_inline(blocks[0].text).strip().casefold() == title.strip().casefold():
            return blocks[1:]
    return blocks


def inline_html(text: str) -> str:
    """Escape first, then add markup. Never the other way around."""
    escaped = html_escape.escape(text)
    escaped = _CODE.sub(r"<code>\1</code>", escaped)
    escaped = _BOLD.sub(r"<strong>\1</strong>", escaped)
    return _ITALIC.sub(r"<em>\1</em>", escaped)


def strip_inline(text: str) -> str:
    """Plain text with the Markdown punctuation removed."""
    plain = _CODE.sub(r"\1", text)
    plain = _BOLD.sub(r"\1", plain)
    return _ITALIC.sub(r"\1", plain)


def backend_for(fmt: str) -> tuple[bool, str]:
    """Is this format usable here, and if not, how is it installed?"""
    if fmt not in FORMATS:
        return False, f"{fmt} is not a format I can write."
    _, package, hint = FORMATS[fmt]
    if package is None:
        return True, ""
    try:
        __import__(package)
    except ImportError:
        return False, f"{fmt.upper()} needs a library that is not installed: {hint}"
    return True, ""


def available_formats() -> list[str]:
    return [fmt for fmt in FORMATS if backend_for(fmt)[0]]


# ---------------------------------------------------------------- documents

def render_txt(blocks: list[Block], title: str) -> bytes:
    lines = [title, "=" * len(title), ""]
    for block in blocks:
        plain = strip_inline(block.text)
        if block.kind == "heading":
            lines += ["", plain, "-" * len(plain)]
        elif block.kind == "bullet":
            lines.append(f"  • {plain}")
        elif block.kind == "number":
            lines.append(f"  - {plain}")
        elif block.kind == "code":
            lines += ["", *(f"    {ln}" for ln in block.text.split("\n")), ""]
        else:
            lines += [plain, ""]
    return ("\n".join(lines).rstrip() + "\n").encode("utf-8")


def render_md(blocks: list[Block], title: str) -> bytes:
    lines = [f"# {title}", ""]
    for block in blocks:
        if block.kind == "heading":
            lines += [f"{'#' * min(block.level + 1, 6)} {block.text}", ""]
        elif block.kind == "bullet":
            lines.append(f"- {block.text}")
        elif block.kind == "number":
            lines.append(f"1. {block.text}")
        elif block.kind == "code":
            lines += ["```", block.text, "```", ""]
        else:
            lines += [block.text, ""]
    return ("\n".join(lines).rstrip() + "\n").encode("utf-8")


def render_html(blocks: list[Block], title: str) -> bytes:
    body: list[str] = []
    listing: str | None = None

    def close() -> None:
        nonlocal listing
        if listing:
            body.append(f"</{listing}>")
            listing = None

    for block in blocks:
        if block.kind in {"bullet", "number"}:
            wanted = "ul" if block.kind == "bullet" else "ol"
            if listing != wanted:
                close()
                body.append(f"<{wanted}>")
                listing = wanted
            body.append(f"<li>{inline_html(block.text)}</li>")
            continue
        close()
        if block.kind == "heading":
            level = min(block.level + 1, 6)
            body.append(f"<h{level}>{inline_html(block.text)}</h{level}>")
        elif block.kind == "code":
            body.append(f"<pre><code>{html_escape.escape(block.text)}</code></pre>")
        else:
            body.append(f"<p>{inline_html(block.text)}</p>")
    close()

    # Self-contained and script-free: this file may be opened from anywhere,
    # and a document the agent wrote should never execute anything.
    return f"""<!doctype html>
<html lang="en">
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{html_escape.escape(title)}</title>
<style>
  body{{max-width:46rem;margin:2rem auto;padding:0 1rem;
       font:16px/1.6 -apple-system,BlinkMacSystemFont,"Segoe UI",Roboto,"Noto Sans",sans-serif;
       color:#1a1a1a;background:#fff}}
  h1,h2,h3{{line-height:1.25}}
  pre{{background:#f4f4f5;padding:.75rem 1rem;border-radius:8px;overflow:auto}}
  code{{font-family:ui-monospace,Menlo,Consolas,monospace;font-size:.9em}}
  @media (prefers-color-scheme:dark){{body{{color:#e8e8ea;background:#141416}}
    pre{{background:#1f1f23}}}}
</style>
<h1>{html_escape.escape(title)}</h1>
{chr(10).join(body)}
""".encode()


def render_pdf(blocks: list[Block], title: str) -> bytes:
    """PDF through reportlab's platypus flowables."""
    try:
        from reportlab.lib.enums import TA_LEFT
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer
    except ImportError as exc:
        raise RenderError(
            "PDF needs reportlab, which is not installed: pip install reportlab",
            code="backend_missing",
        ) from exc

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer, pagesize=A4, title=title,
        leftMargin=20 * mm, rightMargin=20 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm,
    )
    styles = getSampleStyleSheet()
    mono = ParagraphStyle(
        "AgentCode", parent=styles["Code"], fontName="Courier",
        fontSize=9, leading=12, alignment=TA_LEFT,
    )

    story = [Paragraph(inline_html(title), styles["Title"]), Spacer(1, 6)]
    pending: list = []

    def flush_list() -> None:
        if pending:
            story.append(ListFlowable(list(pending), bulletType="bullet", leftIndent=14))
            story.append(Spacer(1, 4))
            pending.clear()

    for block in blocks:
        if block.kind in {"bullet", "number"}:
            pending.append(ListItem(Paragraph(inline_html(block.text), styles["BodyText"])))
            continue
        flush_list()
        if block.kind == "heading":
            style = styles[f"Heading{min(block.level, 4)}"]
            story.append(Paragraph(inline_html(block.text), style))
        elif block.kind == "code":
            for line in block.text.split("\n"):
                story.append(Paragraph(html_escape.escape(line) or "&nbsp;", mono))
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(inline_html(block.text), styles["BodyText"]))
            story.append(Spacer(1, 4))
    flush_list()

    document.build(story)
    return buffer.getvalue()


# ------------------------------------------------------------------- sheets

def _rows_as_text(rows: list[list]) -> list[list[str]]:
    return [["" if cell is None else str(cell) for cell in row] for row in rows]


def render_csv(rows: list[list], title: str) -> bytes:
    del title
    buffer = io.StringIO()
    # QUOTE_MINIMAL with \r\n is what spreadsheet software expects, and
    # utf-8-sig so Excel opens non-ASCII correctly instead of showing mojibake.
    writer = csv.writer(buffer, lineterminator="\r\n")
    writer.writerows(_rows_as_text(rows))
    return buffer.getvalue().encode("utf-8-sig")


def render_json(rows: list[list], title: str) -> bytes:
    """Rows as objects when there is a header, else as a list of lists."""
    if rows and len(rows) > 1:
        header = [str(cell) for cell in rows[0]]
        records = [dict(zip(header, row, strict=False)) for row in rows[1:]]
        payload = {"title": title, "records": records}
    else:
        payload = {"title": title, "rows": rows}
    return (json.dumps(payload, ensure_ascii=False, indent=2) + "\n").encode("utf-8")


def render_xlsx(rows: list[list], title: str) -> bytes:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
    except ImportError as exc:
        raise RenderError(
            "XLSX needs openpyxl, which is not installed: pip install openpyxl",
            code="backend_missing",
        ) from exc

    workbook = Workbook()
    sheet = workbook.active
    # Excel rejects some characters and caps the name at 31 chars.
    sheet.title = (re.sub(r"[\\/*?:\[\]]", " ", title) or "Sheet")[:31]
    for row in rows:
        sheet.append(["" if cell is None else cell for cell in row])
    if rows:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        for index, _ in enumerate(rows[0], start=1):
            column = sheet.cell(row=1, column=index).column_letter
            widest = max(
                (len(str(r[index - 1])) for r in rows if len(r) >= index), default=10
            )
            sheet.column_dimensions[column].width = min(max(widest + 2, 10), 60)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


# ------------------------------------------------------------------- slides

def render_pptx(slides: list[dict], title: str) -> bytes:
    try:
        from pptx import Presentation
        from pptx.util import Inches
    except ImportError as exc:
        raise RenderError(
            "PPTX needs python-pptx, which is not installed. On Termux run "
            "`pkg install python-lxml python-pillow` first, then "
            "`pip install python-pptx`.",
            code="backend_missing",
        ) from exc

    presentation = Presentation()
    cover = presentation.slides.add_slide(presentation.slide_layouts[0])
    cover.shapes.title.text = title

    for slide in slides:
        layout = presentation.slide_layouts[1]
        rendered = presentation.slides.add_slide(layout)
        rendered.shapes.title.text = str(slide.get("title", ""))
        body = rendered.placeholders[1].text_frame
        bullets = slide.get("bullets") or []
        for index, bullet in enumerate(bullets):
            paragraph = body.paragraphs[0] if index == 0 else body.add_paragraph()
            paragraph.text = str(bullet)
        if not bullets:
            body.text = ""
        rendered.shapes.title.width = Inches(9)

    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def render(kind: str, fmt: str, *, title: str, source) -> bytes:
    """Render a stored source into one format."""
    if fmt not in FORMATS:
        raise RenderError(f"{fmt} is not a format I can write.", code="unknown_format")
    expected_kind, _, _ = FORMATS[fmt]
    if expected_kind != kind:
        raise RenderError(
            f"{fmt} is a {expected_kind} format; this is a {kind}.",
            code="format_mismatch",
        )
    usable, reason = backend_for(fmt)
    if not usable:
        raise RenderError(reason, code="backend_missing")

    if kind == DOCUMENT:
        blocks = drop_duplicate_title(parse_markdown(str(source)), title)
        return {
            "md": render_md, "txt": render_txt,
            "html": render_html, "pdf": render_pdf,
        }[fmt](blocks, title)
    if kind == SHEET:
        rows = [list(row) for row in source]
        return {"csv": render_csv, "json": render_json, "xlsx": render_xlsx}[fmt](rows, title)
    return render_pptx(list(source), title)


def read_file_text(path: Path, *, limit: int = 20000) -> str:
    """Extract readable text from a file the agent may not have written."""
    suffix = path.suffix.lower().lstrip(".")
    if suffix == "pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise RenderError(
                "Reading PDF needs pypdf, which is not installed: pip install pypdf",
                code="backend_missing",
            ) from exc
        reader = PdfReader(str(path))
        pages = [page.extract_text() or "" for page in reader.pages]
        return "\n\n".join(pages)[:limit]
    if suffix == "xlsx":
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise RenderError(
                "Reading XLSX needs openpyxl: pip install openpyxl",
                code="backend_missing",
            ) from exc
        workbook = load_workbook(str(path), read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets:
            lines.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                lines.append(
                    ", ".join("" if cell is None else str(cell) for cell in row)
                )
        return "\n".join(lines)[:limit]
    if suffix == "pptx":
        try:
            from pptx import Presentation
        except ImportError as exc:
            raise RenderError(
                "Reading PPTX needs python-pptx: pip install python-pptx",
                code="backend_missing",
            ) from exc
        presentation = Presentation(str(path))
        lines = []
        for number, slide in enumerate(presentation.slides, start=1):
            lines.append(f"## Slide {number}")
            for shape in slide.shapes:
                if shape.has_text_frame and shape.text_frame.text.strip():
                    lines.append(shape.text_frame.text)
        return "\n".join(lines)[:limit]
    # Everything else is text, read defensively.
    return path.read_text(encoding="utf-8", errors="replace")[:limit]
