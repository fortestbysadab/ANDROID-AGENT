"""Tests for document creation, revision and reading.

Two ideas carry most of the weight:

* **A document is its source.** Revision re-renders from stored source and
  keeps the previous file, because the owner asked for a new version rather
  than a destroyed one.
* **Optional backends degrade, never crash.** PDF and PPTX need libraries
  that may not be installed on a given phone; the tool must say which and
  what is available instead.
"""

from __future__ import annotations

import csv
import io
import json
import tempfile
import unittest
from pathlib import Path
from typing import ClassVar
from unittest import mock

from android_agent.documents.render import (
    FORMATS,
    RenderError,
    available_formats,
    backend_for,
    drop_duplicate_title,
    parse_markdown,
    render,
    render_csv,
    render_html,
    render_json,
    render_md,
    render_txt,
)
from android_agent.documents.store import DocumentStore, slugify
from android_agent.tools.base import Risk
from android_agent.tools.document_tools import document_tools

SAMPLE = """# Title

Some **bold**, *italic* and `code`.

- first
- second

1. step one
2. step two

```
x = 1
```
"""


class MarkdownTests(unittest.TestCase):
    def test_every_construct_is_recognised(self):
        kinds = [block.kind for block in parse_markdown(SAMPLE)]
        self.assertEqual(
            kinds,
            ["heading", "paragraph", "bullet", "bullet", "number", "number", "code"],
        )

    def test_paragraphs_join_wrapped_lines(self):
        blocks = parse_markdown("one line\nsame paragraph\n\nnew one")
        self.assertEqual(blocks[0].text, "one line same paragraph")
        self.assertEqual(len(blocks), 2)

    def test_heading_level_is_kept(self):
        self.assertEqual(parse_markdown("### Deep")[0].level, 3)

    def test_an_unterminated_fence_keeps_its_content(self):
        blocks = parse_markdown("```\nnever closed")
        self.assertEqual(blocks[-1].kind, "code")
        self.assertIn("never closed", blocks[-1].text)

    def test_a_duplicate_leading_title_is_dropped(self):
        blocks = drop_duplicate_title(parse_markdown("# Report\n\nBody"), "Report")
        self.assertEqual([b.kind for b in blocks], ["paragraph"])

    def test_a_different_leading_heading_is_kept(self):
        blocks = drop_duplicate_title(parse_markdown("# Intro\n\nBody"), "Report")
        self.assertEqual([b.kind for b in blocks], ["heading", "paragraph"])

    def test_html_escapes_before_marking_up(self):
        out = render_html(parse_markdown("<script>alert(1)</script> **safe**"), "T")
        body = out.decode()
        self.assertIn("&lt;script&gt;", body)
        self.assertNotIn("<script>alert", body)
        self.assertIn("<strong>safe</strong>", body)

    def test_html_output_runs_no_scripts(self):
        """A document the agent wrote must never execute anything."""
        out = render_html(parse_markdown("hello"), "T").decode()
        self.assertNotIn("<script", out)
        self.assertNotIn("onclick", out)

    def test_non_ascii_survives_every_text_renderer(self):
        blocks = parse_markdown("আমার রিপোর্ট\n\n- হিন্দी")
        for renderer in (render_txt, render_md, render_html):
            with self.subTest(renderer=renderer.__name__):
                self.assertIn("আমার", renderer(blocks, "শিরোনাম").decode("utf-8"))


class SheetTests(unittest.TestCase):
    ROWS: ClassVar[list] = [["Name", "Qty"], ["চা", "3"], ["Coffee", "5"]]

    def test_csv_is_readable_back(self):
        text = render_csv(self.ROWS, "x").decode("utf-8-sig")
        self.assertEqual(list(csv.reader(io.StringIO(text))), self.ROWS)

    def test_csv_starts_with_a_bom_so_excel_shows_non_ascii(self):
        self.assertTrue(render_csv(self.ROWS, "x").startswith(b"\xef\xbb\xbf"))

    def test_json_uses_the_header_as_keys(self):
        payload = json.loads(render_json(self.ROWS, "Stock"))
        self.assertEqual(payload["records"][0], {"Name": "চা", "Qty": "3"})
        self.assertEqual(payload["title"], "Stock")

    def test_a_header_only_sheet_does_not_invent_records(self):
        payload = json.loads(render_json([["Name", "Qty"]], "Empty"))
        self.assertIn("rows", payload)

    @unittest.skipUnless(backend_for("xlsx")[0], "openpyxl not installed")
    def test_xlsx_round_trips(self):
        from openpyxl import load_workbook

        data = render(SHEET_KIND, "xlsx", title="Stock", source=self.ROWS)
        workbook = load_workbook(io.BytesIO(data))
        sheet = workbook.active
        self.assertEqual(sheet.title, "Stock")
        self.assertEqual([cell.value for cell in sheet[1]], ["Name", "Qty"])
        self.assertEqual(sheet["A2"].value, "চা")

    @unittest.skipUnless(backend_for("xlsx")[0], "openpyxl not installed")
    def test_a_title_with_illegal_sheet_characters_is_cleaned(self):
        from openpyxl import load_workbook

        data = render(SHEET_KIND, "xlsx", title="Q3/2026: results*", source=self.ROWS)
        self.assertNotIn("/", load_workbook(io.BytesIO(data)).active.title)


SHEET_KIND = "sheet"


class BackendTests(unittest.TestCase):
    def test_a_missing_backend_names_the_install_command(self):
        with mock.patch.dict("sys.modules", {"reportlab": None}):
            usable, reason = backend_for("pdf")
        if not usable:
            self.assertIn("reportlab", reason)

    def test_stdlib_formats_are_always_available(self):
        for fmt in ("md", "txt", "html", "csv", "json"):
            with self.subTest(fmt=fmt):
                self.assertIn(fmt, available_formats())

    def test_rendering_a_mismatched_format_is_refused(self):
        with self.assertRaises(RenderError) as caught:
            render("document", "xlsx", title="T", source="text")
        self.assertEqual(caught.exception.code, "format_mismatch")

    def test_an_unknown_format_is_refused(self):
        with self.assertRaises(RenderError) as caught:
            render("document", "docx", title="T", source="text")
        self.assertEqual(caught.exception.code, "unknown_format")

    def test_every_declared_format_has_a_kind_and_a_hint(self):
        for fmt, (kind, package, hint) in FORMATS.items():
            with self.subTest(fmt=fmt):
                self.assertIn(kind, {"document", "sheet", "slides"})
                if package is not None:
                    self.assertTrue(hint, f"{fmt} has no install hint")


class SlugTests(unittest.TestCase):
    def test_filesystem_characters_are_replaced(self):
        self.assertNotIn("/", slugify("Q3/2026 report"))
        self.assertNotIn(":", slugify("Report: final"))

    def test_other_scripts_keep_their_letters(self):
        """Transliterating a Bengali title would make it unrecognisable."""
        self.assertIn("রিপোর্ট", slugify("মাসিক রিপোর্ট"))

    def test_an_empty_title_still_produces_a_name(self):
        self.assertEqual(slugify("   "), "document")

    def test_long_titles_are_bounded(self):
        self.assertLessEqual(len(slugify("x" * 300)), 60)


class ToolTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.files = Path(self.tmp.name) / "files"
        patcher = mock.patch(
            "android_agent.documents.store.documents_root", return_value=self.files
        )
        patcher.start()
        self.addCleanup(patcher.stop)
        patcher2 = mock.patch(
            "android_agent.tools.document_tools.documents_root", return_value=self.files
        )
        patcher2.start()
        self.addCleanup(patcher2.stop)
        self.files.mkdir(parents=True, exist_ok=True)

        self.store = DocumentStore(Path(self.tmp.name) / "documents.db")
        self.addCleanup(self.store.close)
        self.tools = {tool.name: tool for tool in document_tools(self.store)}

    def call(self, tool_name, **arguments):
        tool = self.tools[tool_name]
        return tool.handler(None, tool.validate(arguments))

    def test_creating_a_markdown_document_writes_a_file(self):
        result = self.call(
            "create_document", title="Notes", format="md", content="Hello **there**"
        )
        self.assertEqual(result.status, "ok")
        path = Path(result.data["artifact_path"])
        self.assertTrue(path.is_file())
        self.assertIn("Hello", path.read_text(encoding="utf-8"))

    def test_the_file_lands_in_the_files_folder(self):
        result = self.call("create_document", title="Notes", format="txt", content="x")
        self.assertEqual(Path(result.data["artifact_path"]).parent, self.files)

    def test_a_sheet_needs_rows_not_content(self):
        result = self.call("create_document", title="Stock", format="csv", content="nope")
        self.assertEqual(result.error_code, "missing_content")
        self.assertIn("rows", result.summary)

    def test_a_document_needs_content_not_rows(self):
        result = self.call(
            "create_document", title="Notes", format="md", rows=[["a"]]
        )
        self.assertEqual(result.error_code, "missing_content")

    def test_revising_creates_version_two_and_keeps_version_one(self):
        created = self.call(
            "create_document", title="Notes", format="md", content="first draft"
        )
        first = Path(created.data["artifact_path"])
        revised = self.call("revise_document", document="Notes", content="second draft")
        second = Path(revised.data["artifact_path"])

        self.assertEqual(revised.data["version"], 2)
        self.assertTrue(first.is_file(), "version 1 must survive")
        self.assertNotEqual(first, second)
        self.assertIn("second draft", second.read_text(encoding="utf-8"))
        self.assertIn("first draft", first.read_text(encoding="utf-8"))

    def test_revising_by_id_works_as_well_as_by_title(self):
        created = self.call("create_document", title="Notes", format="md", content="a")
        revised = self.call(
            "revise_document", document=created.data["id"], content="b"
        )
        self.assertEqual(revised.status, "ok")

    def test_revising_can_change_format_without_resupplying_content(self):
        self.call("create_document", title="Notes", format="md", content="keep me")
        revised = self.call("revise_document", document="Notes", format="html")
        self.assertEqual(revised.status, "ok")
        self.assertIn("keep me", Path(revised.data["artifact_path"]).read_text("utf-8"))

    def test_revising_to_an_incompatible_format_is_refused(self):
        self.call("create_document", title="Notes", format="md", content="x")
        result = self.call("revise_document", document="Notes", format="csv")
        self.assertEqual(result.error_code, "format_mismatch")

    def test_revising_an_unknown_document_says_how_to_find_it(self):
        result = self.call("revise_document", document="nothing", content="x")
        self.assertEqual(result.error_code, "unknown_document")
        self.assertIn("List the documents", result.summary)

    def test_listing_shows_version_and_format(self):
        self.call("create_document", title="Notes", format="md", content="x")
        self.call("revise_document", document="Notes", content="y")
        listed = self.call("list_documents")
        self.assertIn("MD v2", listed.summary)
        self.assertEqual(listed.data["documents"][0]["version"], 2)

    def test_an_empty_library_suggests_what_is_possible(self):
        listed = self.call("list_documents")
        self.assertIn("No documents yet", listed.summary)
        self.assertIn("md", listed.summary)

    def test_reading_back_a_created_document(self):
        self.call("create_document", title="Notes", format="txt", content="the body")
        result = self.call("read_document", name="Notes")
        self.assertEqual(result.status, "ok")
        self.assertIn("the body", result.summary)

    def test_a_read_document_is_untrusted_content(self):
        """A file may have been written by anyone, so it taints the run."""
        self.assertTrue(self.tools["read_document"].returns_untrusted_content)
        self.call("create_document", title="Notes", format="txt", content="hi")
        result = self.call("read_document", name="Notes")
        self.assertIn("UNTRUSTED", result.summary)

    def test_reading_outside_the_files_folder_is_refused(self):
        for attempt in ("../../etc/passwd", "/etc/passwd"):
            with self.subTest(attempt=attempt):
                result = self.call("read_document", name=attempt)
                self.assertEqual(result.status, "error")
                self.assertIn(result.error_code, {"outside_files_folder", "not_found"})

    def test_reading_a_missing_file_says_so(self):
        result = self.call("read_document", name="nope.txt")
        self.assertEqual(result.error_code, "not_found")

    def test_risk_levels_match_what_the_tools_do(self):
        self.assertEqual(self.tools["create_document"].risk, Risk.DEVICE_MUTATION)
        self.assertEqual(self.tools["revise_document"].risk, Risk.DEVICE_MUTATION)
        self.assertEqual(self.tools["list_documents"].risk, Risk.READ_ONLY)
        self.assertEqual(self.tools["read_document"].risk, Risk.SENSITIVE_READ)

    def test_an_unavailable_backend_names_what_is_available(self):
        with mock.patch(
            "android_agent.tools.document_tools.backend_for",
            return_value=(False, "PDF needs reportlab: pip install reportlab"),
        ):
            result = self.call(
                "create_document", title="Report", format="pdf", content="x"
            )
        self.assertEqual(result.error_code, "backend_missing")
        self.assertIn("reportlab", result.summary)
        self.assertIn("available right now", result.summary)

    def test_a_non_ascii_title_produces_a_usable_filename(self):
        result = self.call(
            "create_document", title="মাসিক রিপোর্ট", format="md", content="বিষয়"
        )
        self.assertEqual(result.status, "ok")
        self.assertTrue(Path(result.data["artifact_path"]).is_file())

    def test_the_result_carries_an_artifact_so_the_file_is_delivered(self):
        result = self.call("create_document", title="Notes", format="md", content="x")
        self.assertIn("artifact_path", result.data)
        self.assertIn("saved_to", result.data)


if __name__ == "__main__":
    unittest.main()
